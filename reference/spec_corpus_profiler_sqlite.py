from __future__ import annotations

"""
Word Specification Corpus Profiler - SQLite Edition
====================================================

Scans all .docx files recursively beneath the current Windows user's
Desktop\\Specs folder and builds a structural corpus profile in SQLite.

No --root argument is required for normal use. The default source is discovered
from the current user's Windows Desktop, including redirected/OneDrive Desktop
locations when Windows reports one.

Default outputs:
    Desktop\\Specs\\_Corpus_Profile\\spec_corpus_profile.db
    Desktop\\Specs\\_Corpus_Profile\\corpus_summary.csv
    Desktop\\Specs\\_Corpus_Profile\\layout_families.csv
    Desktop\\Specs\\_Corpus_Profile\\variant_layouts.csv
    Desktop\\Specs\\_Corpus_Profile\\table_signatures.csv
    Desktop\\Specs\\_Corpus_Profile\\table_variants.csv
    Desktop\\Specs\\_Corpus_Profile\\section_coverage.csv
    Desktop\\Specs\\_Corpus_Profile\\heading_variants.csv
    Desktop\\Specs\\_Corpus_Profile\\exception_summary.csv
    Desktop\\Specs\\_Corpus_Profile\\corpus_profile_summary.json
    Desktop\\Specs\\_Corpus_Profile\\spec_corpus_profile_shareable.zip

The profiler is structural: it stores file paths locally for traceability, but
aggregate/shareable exports do not include specification data values or paths.
Source Word files are opened read-only and never modified.
"""

import argparse
import csv
import hashlib
import json
import os
import platform
import re
import shutil
import sqlite3
import sys
import time
import traceback
import zipfile
from collections import Counter
from concurrent.futures import ProcessPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

try:
    from docx import Document
    from docx.document import Document as _Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError as exc:
    raise SystemExit("python-docx is required. Install it with: pip install python-docx") from exc

PROFILER_VERSION = "2.1.0-sqlite"
SCHEMA_VERSION = 4
DEFAULT_DB_FILENAME = "spec_corpus_profile.db"
DEFAULT_OUTPUT_DIRNAME = "_Corpus_Profile"

# -----------------------------------------------------------------------------
# Expected 11-table specification structure
# -----------------------------------------------------------------------------
EXPECTED_TABLES: dict[int, dict[str, Any]] = {
    1: {
        "section": "Product Description",
        "container": "header",
        "aliases": ["product description"],
    },
    2: {
        "section": "Locations",
        "container": "body",
        "aliases": ["locations", "location"],
    },
    3: {
        "section": "Bill of Materials",
        "container": "body",
        "aliases": ["bill of materials", "bom", "bill of material"],
    },
    4: {
        "section": "Secondary Approved Materials",
        "container": "body",
        "aliases": [
            "secondary approved materials",
            "secondary approved material",
            "approved secondary materials",
            "alternate approved materials",
            "alternative approved materials",
        ],
    },
    5: {
        "section": "Process Routing",
        "container": "body",
        "aliases": ["process routing", "routing", "process route"],
    },
    6: {
        "section": "Physical Attributes & Testing",
        "container": "body",
        "aliases": [
            "physical attributes & testing",
            "physical attributes and testing",
            "physical attributes/testing",
            "physical attributes",
        ],
    },
    7: {
        "section": "Slitting Information",
        "container": "body",
        "aliases": ["slitting information", "slit information", "slitting"],
    },
    8: {
        "section": "Packing Information",
        "container": "body",
        "aliases": [
            "packing information",
            "packaging information",
            "packing",
            "packaging",
        ],
    },
    9: {
        "section": "Reporting Requirements",
        "container": "body",
        "aliases": ["reporting requirements", "reporting requirement"],
    },
    10: {
        "section": "Customer Sampling Requirements",
        "container": "body",
        "aliases": [
            "customer sampling requirements",
            "customer sample requirements",
            "sampling requirements",
        ],
    },
    11: {
        "section": "Revision History",
        "container": "body",
        "aliases": ["revision history", "revision histories", "revision log"],
    },
}

ALIAS_TO_NUMBER: dict[str, int] = {}
for _num, _meta in EXPECTED_TABLES.items():
    for _alias in _meta["aliases"]:
        ALIAS_TO_NUMBER[_alias] = _num


# Generic schema/label vocabulary. Only these recognized labels are persisted
# from table cell text. Arbitrary data values are not stored by the profiler.
GENERIC_LABEL_PATTERNS: list[tuple[str, str]] = [
    ("customer location", "customer_location"),
    ("spec originator", "spec_originator"),
    ("phone number", "phone_number"),
    ("structure description", "structure_description"),
    ("date of issue", "date_of_issue"),
    ("structure code", "structure_code"),
    ("revision #", "revision_number"),
    ("revision number", "revision_number"),
    ("spec #", "spec_number"),
    ("spec number", "spec_number"),
    ("raw material item code", "raw_material_item_code"),
    ("raw material", "raw_material"),
    ("basis wt", "basis_weight"),
    ("basis weight", "basis_weight"),
    ("caliper", "caliper"),
    ("supplier", "supplier"),
    ("process conditions", "process_conditions"),
    ("conditions", "conditions"),
    ("comments", "comments"),
    ("machine", "machine"),
    ("process", "process"),
    ("pass", "pass"),
    ("attribute", "attribute"),
    ("coa", "coa"),
    ("units", "units"),
    ("lsl", "lsl"),
    ("lcl", "lcl"),
    ("target", "target"),
    ("ucl", "ucl"),
    ("usl", "usl"),
    ("procedure", "procedure"),
    ("frequency", "frequency"),
    ("slit width", "slit_width"),
    ("width tolerance", "width_tolerance"),
    ("core id", "core_id"),
    ("core type", "core_type"),
    ("splice tape", "splice_tape"),
    ("splice type", "splice_type"),
    ("splice location", "splice_location"),
    ("splice stamp", "splice_stamp_flag"),
    ("cutoff", "cutoff"),
    ("lineal ft", "lineal_ft_per_roll"),
    ("eyemark location", "eyemark_location"),
    ("maximum od", "maximum_od"),
    ("maximum weight", "maximum_weight"),
    ("maximum # splices", "maximum_splices"),
    ("number of cuts", "number_of_cuts"),
    ("unwind", "unwind"),
    ("labels across", "labels_across"),
    ("core tags", "core_tags"),
    ("splice code", "splice_code"),
    ("splicing", "splicing"),
    ("end of roll indicator", "end_of_roll_indicator"),
    ("pallet size", "pallet_size"),
    ("pallet type", "pallet_type"),
    ("rolls / layer", "rolls_per_layer"),
    ("layers / pallet", "layers_per_pallet"),
    ("rolls / pallet", "rolls_per_pallet"),
    ("feet/pallet", "feet_per_pallet"),
    ("poly bags", "poly_bags"),
    ("core plugs", "core_plugs"),
    ("roll cradles", "roll_cradles"),
    ("tier dividers", "tier_dividers"),
    ("top layer", "top_layer"),
    ("end boards", "end_boards"),
    ("pallet labels", "pallet_labels"),
    ("protective corners", "protective_corners"),
    ("stretch wrap pallet", "stretch_wrap_pallet"),
    ("band pallet", "band_pallet"),
    ("wooden top", "wooden_top"),
    ("packing instructions", "packing_instructions"),
    ("additional instructions", "additional_instructions"),
    ("coa required", "coa_required"),
    ("coa recipient", "coa_recipients"),
    ("coa format", "coa_format"),
    ("coa frequency", "coa_frequency"),
    ("values to report on coa", "coa_values"),
    ("samples required", "samples_required"),
    ("send to", "send_to"),
    ("number of samples", "number_of_samples"),
    ("sample preparation", "sample_preparation"),
    ("labeling", "labeling"),
    ("revision", "revision"),
    ("who", "who"),
    ("date", "date"),
    ("facility", "facility"),
    ("customer", "customer"),
    ("item", "item"),
]


# -----------------------------------------------------------------------------
# Utility helpers
# -----------------------------------------------------------------------------
def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def norm(value: str | None) -> str:
    if not value:
        return ""
    value = value.replace("\r", "\n")
    value = re.sub(r"\s+", " ", value).strip().lower()
    value = value.replace("–", "-").replace("—", "-")
    value = value.replace("＆", "&")
    return value


def safe_text(value: str | None, limit: int = 250) -> str:
    text = re.sub(r"\s+", " ", (value or "").replace("\r", " ").replace("\n", " ")).strip()
    return text[:limit]


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def stable_hash(text: str, length: int = 24) -> str:
    return hashlib.sha256(text.encode("utf-8", errors="replace")).hexdigest()[:length]


def document_key(relative_path: str) -> str:
    return "DOC_" + stable_hash(relative_path, 16).upper()


def is_heading_paragraph(p: Paragraph) -> bool:
    style_name = (p.style.name or "") if p.style is not None else ""
    return bool(re.match(r"^heading\s*\d*", style_name.strip(), flags=re.I))


def detect_expected_number(heading: str | None) -> int:
    n = norm(heading)
    if not n:
        return 0
    if n in ALIAS_TO_NUMBER:
        return ALIAS_TO_NUMBER[n]
    # Conservative normalization for punctuation/ampersand differences.
    n2 = n.replace(" and ", " & ")
    n2 = re.sub(r"[^a-z0-9#&]+", " ", n2)
    n2 = re.sub(r"\s+", " ", n2).strip()
    for alias, number in ALIAS_TO_NUMBER.items():
        a2 = alias.replace(" and ", " & ")
        a2 = re.sub(r"[^a-z0-9#&]+", " ", a2)
        a2 = re.sub(r"\s+", " ", a2).strip()
        if n2 == a2:
            return number
    return 0


def section_name(number: int) -> str:
    return EXPECTED_TABLES.get(number, {}).get("section", "Unmapped")


def iter_body_blocks(parent: _Document) -> Iterable[Paragraph | Table]:
    body = parent.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def cell_texts(table: Table) -> Iterable[str]:
    for row in table.rows:
        for cell in row.cells:
            yield cell.text or ""


def generic_labels(table: Table) -> list[str]:
    found: set[str] = set()
    # Match labels across the whole table, but persist only generic vocabulary.
    for text in cell_texts(table):
        n = norm(text)
        if not n:
            continue
        for phrase, token in GENERIC_LABEL_PATTERNS:
            # Labels often include punctuation or a value in the same merged cell.
            if n == phrase or n.startswith(phrase + ":") or n.startswith(phrase + " "):
                found.add(token)
    return sorted(found)


def generic_token_for_cell(value: str) -> str:
    n = norm(value)
    if not n:
        return ""
    for phrase, token in GENERIC_LABEL_PATTERNS:
        if n == phrase or n.startswith(phrase + ":") or n.startswith(phrase + " "):
            return token
    return "_"


def classify_cell(value: str) -> str:
    v = safe_text(value, 500)
    if not v:
        return "E"  # empty
    n = norm(v)
    for phrase, _token in GENERIC_LABEL_PATTERNS:
        if n == phrase or n.startswith(phrase + ":"):
            return "L"  # recognized label
    if re.fullmatch(r"[-+]?\d+(?:[.,]\d+)?(?:\s*[%\"])?", v):
        return "N"
    if re.fullmatch(r"\d{1,2}[/-]\d{1,2}[/-]\d{2,4}", v):
        return "D"
    if "\n" in value:
        return "M"  # multiline
    if len(v) <= 20:
        return "S"
    if len(v) <= 80:
        return "T"
    return "X"  # long text


def table_geometry(table: Table) -> dict[str, Any]:
    rows = table.rows
    row_widths = [len(row.cells) for row in rows]
    physical_row_widths = [len(row._tr.tc_lst) for row in rows]

    hmerge = 0
    vmerge = 0
    physical_cells = 0
    logical_cells = sum(row_widths)
    first_row_merge_patterns: list[str] = []

    for row_idx, row in enumerate(rows):
        row_hmerge = 0
        row_vmerge = 0
        for tc in row._tr.tc_lst:
            physical_cells += 1
            tcPr = tc.tcPr
            if tcPr is None:
                continue
            grid_span = tcPr.gridSpan
            if grid_span is not None:
                try:
                    if int(grid_span.val) > 1:
                        hmerge += 1
                        row_hmerge += 1
                except Exception:
                    pass
            if tcPr.vMerge is not None:
                vmerge += 1
                row_vmerge += 1
        if row_idx < 4:
            first_row_merge_patterns.append(f"h{row_hmerge}v{row_vmerge}")

    # Diagnostic shape only. It is intentionally NOT used for schema grouping
    # because data values can change S/T/N classifications between documents.
    shape_rows: list[str] = []
    for row in rows[:4]:
        shape_rows.append("".join(classify_cell(cell.text or "") for cell in row.cells))

    # Generic token patterns preserve labels but replace arbitrary values with _.
    token_rows: list[str] = []
    for row in rows[:2]:
        token_rows.append("|".join(generic_token_for_cell(cell.text or "") for cell in row.cells))

    try:
        nested_tables = max(0, len(table._tbl.xpath(".//w:tbl")) - 1)
    except Exception:
        nested_tables = 0

    return {
        "row_count": len(rows),
        "min_grid_columns": min(row_widths) if row_widths else 0,
        "max_grid_columns": max(row_widths) if row_widths else 0,
        "row_width_pattern": ",".join(map(str, row_widths[:8])),
        "physical_row_width_pattern": ",".join(map(str, physical_row_widths[:8])),
        "first_row_merge_pattern": "/".join(first_row_merge_patterns),
        "generic_token_pattern": "/".join(token_rows),
        "physical_cell_count": physical_cells,
        "logical_cell_count": logical_cells,
        "horizontal_merge_count": hmerge,
        "vertical_merge_count": vmerge,
        "nested_table_count": nested_tables,
        "cell_shape": "/".join(shape_rows),
    }


def table_structural_signature(
    table: Table,
    container_type: str,
    expected_number: int,
    preceding_heading: str,
) -> tuple[str, str, dict[str, Any]]:
    geom = table_geometry(table)
    labels = generic_labels(table)
    container_group = (
        "header" if "header" in container_type
        else "footer" if "footer" in container_type
        else "body"
    )

    # Broad schema signature: semantic structure, intentionally tolerant of
    # row counts, value lengths, and Word XML merge oddities. For known T01-T11
    # sections, the expected section + generic label vocabulary is the core.
    schema_payload = {
        "container": container_group,
        "expected_number": expected_number,
        "heading": (
            section_name(expected_number).lower()
            if expected_number > 0
            else norm(preceding_heading)
        ),
        "labels": labels,
    }
    # Unknown/unmapped tables need extra structural detail to avoid collapsing
    # unrelated tables together. Known T01-T11 schemas are grouped semantically;
    # their Word-grid differences are captured by the strict variant signature.
    if expected_number == 0:
        schema_payload["generic_token_pattern"] = geom["generic_token_pattern"]
        schema_payload["max_cols"] = geom["max_grid_columns"]
        schema_payload["min_cols"] = geom["min_grid_columns"]

    schema_encoded = json.dumps(schema_payload, sort_keys=True, separators=(",", ":"))
    schema_signature = "TS_" + stable_hash(schema_encoded, 20).upper()

    # Strict variant signature: captures XML/grid differences that may require
    # different parser logic even when the semantic table is the same.
    variant_payload = {
        **schema_payload,
        "min_cols": geom["min_grid_columns"],
        "max_cols": geom["max_grid_columns"],
        "first_row_widths": geom["row_width_pattern"].split(",")[:4],
        "first_physical_widths": geom["physical_row_width_pattern"].split(",")[:4],
        "first_merge_pattern": geom["first_row_merge_pattern"],
        "nested": geom["nested_table_count"],
    }
    variant_encoded = json.dumps(variant_payload, sort_keys=True, separators=(",", ":"))
    variant_signature = "TV_" + stable_hash(variant_encoded, 20).upper()

    geom["generic_label_tokens"] = ",".join(labels)
    return schema_signature, variant_signature, geom


def table_xml_hash(table: Table) -> str:
    # Used only to deduplicate identical linked headers/footers inside one doc.
    return hashlib.sha1(table._tbl.xml.encode("utf-8", errors="replace")).hexdigest()


def profile_table(
    table: Table,
    container_type: str,
    container_index: int,
    ordinal: int,
    source_order: int,
    preceding_heading: str = "",
    heading_style: str = "",
    forced_expected_number: int = 0,
) -> dict[str, Any]:
    expected = forced_expected_number or detect_expected_number(preceding_heading)
    signature, variant_signature, geom = table_structural_signature(
        table, container_type, expected, preceding_heading
    )
    return {
        "container_type": container_type,
        "container_index": container_index,
        "table_ordinal": ordinal,
        "source_order": source_order,
        "expected_table_number": expected,
        "detected_section": section_name(expected),
        "preceding_heading": safe_text(preceding_heading, 180),
        "heading_style": safe_text(heading_style, 80),
        "table_signature": signature,
        "table_variant_signature": variant_signature,
        **geom,
    }


def header_expected_number(table: Table) -> int:
    labels = set(generic_labels(table))
    core = {"customer", "spec_number", "item"}
    if len(core.intersection(labels)) >= 2:
        return 1
    # Some templates merge labels and values in a way that defeats exact label
    # recognition. Product-description header tables are still commonly rich in
    # these labels; use a conservative threshold.
    secondary = {
        "structure_description",
        "revision_number",
        "structure_code",
        "date_of_issue",
    }
    if "item" in labels and len(secondary.intersection(labels)) >= 2:
        return 1
    return 0


def profile_one_file(task: tuple[str, str, int, int]) -> dict[str, Any]:
    """Worker-safe function. Returns structural information only."""
    abs_path_str, relative_path, file_size, modified_ns = task
    path = Path(abs_path_str)
    started = time.perf_counter()

    result: dict[str, Any] = {
        "relative_path": relative_path,
        "document_key": document_key(relative_path),
        "filename": path.name,
        "parent_folder": str(Path(relative_path).parent),
        "file_size_bytes": file_size,
        "modified_ns": modified_ns,
        "status": "error",
        "error_type": "",
        "error_message": "",
        "traceback": "",
        "sha256": "",
        "paragraph_count": 0,
        "heading_count": 0,
        "body_table_count": 0,
        "header_table_count": 0,
        "footer_table_count": 0,
        "total_table_count": 0,
        "known_sections_found": 0,
        "unknown_heading_count": 0,
        "layout_signature": "",
        "layout_sequence": "[]",
        "variant_layout_signature": "",
        "variant_layout_sequence": "[]",
        "section_number_sequence": "",
        "headings": [],
        "tables": [],
        "elapsed_seconds": 0.0,
    }

    try:
        result["sha256"] = sha256_file(path)
        doc = Document(str(path))
        result["paragraph_count"] = len(doc.paragraphs)

        headings: list[dict[str, str]] = []
        for p in doc.paragraphs:
            text = safe_text(p.text, 180)
            if text and is_heading_paragraph(p):
                style = (p.style.name or "") if p.style is not None else ""
                headings.append(
                    {
                        "heading_text": text,
                        "heading_normalized": norm(text),
                        "style_name": style,
                        "expected_table_number": detect_expected_number(text),
                    }
                )
        result["headings"] = headings
        result["heading_count"] = len(headings)
        result["unknown_heading_count"] = sum(
            1 for h in headings if h["expected_table_number"] == 0
        )

        tables: list[dict[str, Any]] = []
        source_order = 0

        # Header variants: primary, first-page, and even-page header across all
        # sections. Linked/repeated header tables are deduplicated by XML hash.
        seen_header_xml: set[str] = set()
        header_ordinal = 0
        for section_idx, section in enumerate(doc.sections, start=1):
            header_variants = [
                ("header", section.header),
                ("first_page_header", section.first_page_header),
                ("even_page_header", section.even_page_header),
            ]
            for container_name, hdr in header_variants:
                for table in hdr.tables:
                    xml_hash = table_xml_hash(table)
                    if xml_hash in seen_header_xml:
                        continue
                    seen_header_xml.add(xml_hash)
                    header_ordinal += 1
                    source_order += 1
                    expected = header_expected_number(table)
                    tables.append(
                        profile_table(
                            table,
                            container_name,
                            section_idx,
                            header_ordinal,
                            source_order,
                            preceding_heading=(
                                "Product Description" if expected == 1 else ""
                            ),
                            heading_style="Header Table",
                            forced_expected_number=expected,
                        )
                    )

        # Body tables in true document order, linked to the nearest preceding
        # Word Heading paragraph.
        current_heading = ""
        current_heading_style = ""
        body_ordinal = 0
        for block in iter_body_blocks(doc):
            if isinstance(block, Paragraph):
                text = safe_text(block.text, 180)
                if text and is_heading_paragraph(block):
                    current_heading = text
                    current_heading_style = (
                        block.style.name if block.style is not None else ""
                    )
            elif isinstance(block, Table):
                body_ordinal += 1
                source_order += 1
                tables.append(
                    profile_table(
                        block,
                        "body",
                        0,
                        body_ordinal,
                        source_order,
                        preceding_heading=current_heading,
                        heading_style=current_heading_style,
                    )
                )

        # Footer variants are structural anomalies for these specs, but we still
        # profile them so old templates cannot hide unexpected tables.
        seen_footer_xml: set[str] = set()
        footer_ordinal = 0
        for section_idx, section in enumerate(doc.sections, start=1):
            footer_variants = [
                ("footer", section.footer),
                ("first_page_footer", section.first_page_footer),
                ("even_page_footer", section.even_page_footer),
            ]
            for container_name, ftr in footer_variants:
                for table in ftr.tables:
                    xml_hash = table_xml_hash(table)
                    if xml_hash in seen_footer_xml:
                        continue
                    seen_footer_xml.add(xml_hash)
                    footer_ordinal += 1
                    source_order += 1
                    tables.append(
                        profile_table(
                            table,
                            container_name,
                            section_idx,
                            footer_ordinal,
                            source_order,
                            preceding_heading="",
                            heading_style="Footer Table",
                        )
                    )

        result["tables"] = tables
        result["body_table_count"] = sum(
            1 for t in tables if t["container_type"] == "body"
        )
        result["header_table_count"] = sum(
            1 for t in tables if "header" in t["container_type"]
        )
        result["footer_table_count"] = sum(
            1 for t in tables if "footer" in t["container_type"]
        )
        result["total_table_count"] = len(tables)

        known_nums = [
            t["expected_table_number"]
            for t in tables
            if t["expected_table_number"] > 0
        ]
        result["known_sections_found"] = len(set(known_nums))

        # Broad layout family groups documents that expose the same semantic
        # table schemas. Variant layout separately captures Word XML/grid quirks.
        layout_items = [
            {
                "container": (
                    "header" if "header" in t["container_type"]
                    else "footer" if "footer" in t["container_type"]
                    else "body"
                ),
                "n": t["expected_table_number"],
                "sig": t["table_signature"],
            }
            for t in tables
        ]
        variant_items = [
            {
                "container": (
                    "header" if "header" in t["container_type"]
                    else "footer" if "footer" in t["container_type"]
                    else "body"
                ),
                "n": t["expected_table_number"],
                "sig": t["table_variant_signature"],
            }
            for t in tables
        ]
        layout_encoded = json.dumps(layout_items, sort_keys=True, separators=(",", ":"))
        variant_encoded = json.dumps(variant_items, sort_keys=True, separators=(",", ":"))
        result["layout_signature"] = "LY_" + stable_hash(layout_encoded, 20).upper()
        result["layout_sequence"] = json.dumps(layout_items, separators=(",", ":"))
        result["variant_layout_signature"] = "VL_" + stable_hash(variant_encoded, 20).upper()
        result["variant_layout_sequence"] = json.dumps(variant_items, separators=(",", ":"))
        result["section_number_sequence"] = ",".join(str(n) for n in known_nums)
        result["status"] = "ok"
    except Exception as exc:
        result["error_type"] = type(exc).__name__
        result["error_message"] = safe_text(str(exc), 1000)
        result["traceback"] = traceback.format_exc(limit=8)
    finally:
        result["elapsed_seconds"] = round(time.perf_counter() - started, 4)

    return result


# -----------------------------------------------------------------------------
# Discovery
# -----------------------------------------------------------------------------
def discover_docx(root: Path, output_dir: Path, exclude_dirs: set[str]) -> list[tuple[Path, str, int, int]]:
    files: list[tuple[Path, str, int, int]] = []
    root_resolved = root.resolve()
    output_resolved = output_dir.resolve()
    exclude_lower = {d.lower() for d in exclude_dirs}

    for dirpath, dirnames, filenames in os.walk(root, followlinks=False):
        current = Path(dirpath)

        # Prevent walking the profiler's own output tree if it sits under root.
        kept_dirs = []
        for d in dirnames:
            candidate = (current / d).resolve()
            if d.lower() in exclude_lower:
                continue
            try:
                candidate.relative_to(output_resolved)
                # candidate is output_dir or inside it
                continue
            except ValueError:
                pass
            kept_dirs.append(d)
        dirnames[:] = kept_dirs

        for filename in filenames:
            if filename.startswith("~$"):
                continue
            if not filename.lower().endswith(".docx"):
                continue
            path = current / filename
            try:
                stat = path.stat()
            except OSError:
                # Let the normal profiler worker record the access failure with
                # placeholder stat information.
                stat_size = -1
                stat_mtime_ns = -1
            else:
                stat_size = stat.st_size
                stat_mtime_ns = stat.st_mtime_ns
            try:
                rel = str(path.resolve().relative_to(root_resolved))
            except Exception:
                rel = str(path.relative_to(root))
            files.append((path, rel, stat_size, stat_mtime_ns))

    files.sort(key=lambda x: x[1].lower())
    return files


# -----------------------------------------------------------------------------
# Windows path discovery
# -----------------------------------------------------------------------------
def find_windows_desktop() -> Path:
    """Return the current user's Windows Desktop, including redirected desktops."""
    if os.name == "nt":
        try:
            import ctypes

            CSIDL_DESKTOPDIRECTORY = 0x0010
            SHGFP_TYPE_CURRENT = 0
            buffer = ctypes.create_unicode_buffer(32768)
            result = ctypes.windll.shell32.SHGetFolderPathW(
                None,
                CSIDL_DESKTOPDIRECTORY,
                None,
                SHGFP_TYPE_CURRENT,
                buffer,
            )
            if result == 0 and buffer.value:
                return Path(buffer.value)
        except Exception:
            pass

    candidates: list[Path] = []
    userprofile = os.environ.get("USERPROFILE")
    onedrive = os.environ.get("OneDrive") or os.environ.get("OneDriveCommercial")
    if userprofile:
        candidates.append(Path(userprofile) / "Desktop")
    if onedrive:
        candidates.append(Path(onedrive) / "Desktop")
    candidates.append(Path.home() / "Desktop")

    for candidate in candidates:
        if candidate.exists():
            return candidate
    return candidates[0] if candidates else Path.home() / "Desktop"


def default_specs_root() -> Path:
    return find_windows_desktop() / "Specs"


# -----------------------------------------------------------------------------
# SQLite persistence
# -----------------------------------------------------------------------------

def connect_db(db_path: Path) -> sqlite3.Connection:
    conn = sqlite3.connect(str(db_path), timeout=60)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA journal_mode = WAL")
    conn.execute("PRAGMA synchronous = NORMAL")
    conn.execute("PRAGMA busy_timeout = 60000")
    return conn


def q_all(conn: sqlite3.Connection, sql: str, params: tuple[Any, ...] = ()) -> list[dict[str, Any]]:
    return [dict(r) for r in conn.execute(sql, params).fetchall()]


def q_one(conn: sqlite3.Connection, sql: str, params: tuple[Any, ...] = ()) -> dict[str, Any] | None:
    row = conn.execute(sql, params).fetchone()
    return dict(row) if row is not None else None


def scalar(conn: sqlite3.Connection, sql: str, params: tuple[Any, ...] = ()) -> Any:
    row = conn.execute(sql, params).fetchone()
    return None if row is None else row[0]


SCHEMA_SQL = r"""
CREATE TABLE IF NOT EXISTS T00_Profile_Runs (
    run_id INTEGER PRIMARY KEY AUTOINCREMENT,
    started_at TEXT NOT NULL,
    completed_at TEXT,
    root_path TEXT NOT NULL,
    output_db TEXT NOT NULL,
    profiler_version TEXT NOT NULL,
    schema_version INTEGER NOT NULL,
    python_version TEXT NOT NULL,
    platform TEXT,
    workers INTEGER NOT NULL,
    full_rescan INTEGER NOT NULL DEFAULT 0,
    files_discovered INTEGER NOT NULL DEFAULT 0,
    files_profiled INTEGER NOT NULL DEFAULT 0,
    files_skipped INTEGER NOT NULL DEFAULT 0,
    files_failed INTEGER NOT NULL DEFAULT 0,
    elapsed_seconds REAL
);

CREATE TABLE IF NOT EXISTS T00_Profile_Documents (
    document_id INTEGER PRIMARY KEY AUTOINCREMENT,
    document_key TEXT NOT NULL UNIQUE,
    relative_path TEXT NOT NULL UNIQUE,
    filename TEXT NOT NULL,
    parent_folder TEXT,
    file_size_bytes INTEGER,
    modified_ns INTEGER,
    sha256 TEXT,
    profiled_at TEXT,
    status TEXT NOT NULL,
    error_type TEXT,
    error_message TEXT,
    paragraph_count INTEGER,
    heading_count INTEGER,
    body_table_count INTEGER,
    header_table_count INTEGER,
    footer_table_count INTEGER,
    total_table_count INTEGER,
    known_sections_found INTEGER,
    unknown_heading_count INTEGER,
    layout_signature TEXT,
    layout_sequence TEXT,
    variant_layout_signature TEXT,
    variant_layout_sequence TEXT,
    section_number_sequence TEXT,
    elapsed_seconds REAL,
    is_present INTEGER NOT NULL DEFAULT 1,
    last_seen_run_id INTEGER,
    FOREIGN KEY(last_seen_run_id) REFERENCES T00_Profile_Runs(run_id)
);

CREATE TABLE IF NOT EXISTS T00_Profile_Heading_Occurrences (
    heading_occurrence_id INTEGER PRIMARY KEY AUTOINCREMENT,
    document_id INTEGER NOT NULL,
    heading_ordinal INTEGER NOT NULL,
    heading_text TEXT NOT NULL,
    heading_normalized TEXT NOT NULL,
    style_name TEXT,
    expected_table_number INTEGER NOT NULL DEFAULT 0,
    FOREIGN KEY(document_id) REFERENCES T00_Profile_Documents(document_id) ON DELETE CASCADE
);

CREATE TABLE IF NOT EXISTS T00_Profile_Table_Occurrences (
    table_occurrence_id INTEGER PRIMARY KEY AUTOINCREMENT,
    document_id INTEGER NOT NULL,
    source_order INTEGER NOT NULL,
    container_type TEXT NOT NULL,
    container_index INTEGER NOT NULL DEFAULT 0,
    table_ordinal INTEGER NOT NULL,
    expected_table_number INTEGER NOT NULL DEFAULT 0,
    detected_section TEXT,
    preceding_heading TEXT,
    heading_style TEXT,
    table_signature TEXT NOT NULL,
    table_variant_signature TEXT NOT NULL,
    row_count INTEGER NOT NULL,
    min_grid_columns INTEGER NOT NULL,
    max_grid_columns INTEGER NOT NULL,
    row_width_pattern TEXT,
    physical_row_width_pattern TEXT,
    first_row_merge_pattern TEXT,
    generic_token_pattern TEXT,
    physical_cell_count INTEGER,
    logical_cell_count INTEGER,
    horizontal_merge_count INTEGER,
    vertical_merge_count INTEGER,
    nested_table_count INTEGER,
    cell_shape TEXT,
    generic_label_tokens TEXT,
    FOREIGN KEY(document_id) REFERENCES T00_Profile_Documents(document_id) ON DELETE CASCADE
);

CREATE TABLE IF NOT EXISTS T00_Profile_Run_Errors (
    error_id INTEGER PRIMARY KEY AUTOINCREMENT,
    run_id INTEGER NOT NULL,
    document_key TEXT,
    relative_path TEXT,
    error_type TEXT,
    error_message TEXT,
    traceback TEXT,
    created_at TEXT NOT NULL,
    FOREIGN KEY(run_id) REFERENCES T00_Profile_Runs(run_id)
);

CREATE TABLE IF NOT EXISTS T00_Profile_Layouts (
    layout_signature TEXT PRIMARY KEY,
    document_count INTEGER NOT NULL,
    percent_of_ok_documents REAL NOT NULL,
    section_number_sequence TEXT,
    layout_sequence TEXT,
    min_total_tables INTEGER,
    max_total_tables INTEGER,
    avg_total_tables REAL,
    min_body_tables INTEGER,
    max_body_tables INTEGER,
    min_header_tables INTEGER,
    max_header_tables INTEGER,
    first_document_key TEXT
);
CREATE TABLE IF NOT EXISTS T00_Profile_Variant_Layouts (
    variant_layout_signature TEXT PRIMARY KEY,
    document_count INTEGER NOT NULL,
    percent_of_ok_documents REAL NOT NULL,
    section_number_sequence TEXT,
    variant_layout_sequence TEXT,
    first_document_key TEXT
);
CREATE TABLE IF NOT EXISTS T00_Profile_Table_Signatures (
    table_signature TEXT PRIMARY KEY,
    expected_table_number INTEGER NOT NULL,
    detected_section TEXT,
    container_type TEXT,
    document_count INTEGER NOT NULL,
    occurrence_count INTEGER NOT NULL,
    percent_of_ok_documents REAL NOT NULL,
    min_row_count INTEGER,
    max_row_count INTEGER,
    avg_row_count REAL,
    min_grid_columns INTEGER,
    max_grid_columns INTEGER,
    row_width_pattern TEXT,
    physical_row_width_pattern TEXT,
    horizontal_merge_count INTEGER,
    vertical_merge_count INTEGER,
    nested_table_count INTEGER,
    cell_shape TEXT,
    generic_label_tokens TEXT
);
CREATE TABLE IF NOT EXISTS T00_Profile_Table_Variants (
    table_variant_signature TEXT PRIMARY KEY,
    table_signature TEXT NOT NULL,
    expected_table_number INTEGER NOT NULL,
    detected_section TEXT,
    container_type TEXT,
    document_count INTEGER NOT NULL,
    occurrence_count INTEGER NOT NULL,
    percent_of_ok_documents REAL NOT NULL,
    min_row_count INTEGER,
    max_row_count INTEGER,
    avg_row_count REAL,
    min_grid_columns INTEGER,
    max_grid_columns INTEGER,
    first_row_merge_pattern TEXT,
    generic_token_pattern TEXT,
    generic_label_tokens TEXT
);
CREATE TABLE IF NOT EXISTS T00_Profile_Heading_Variants (
    heading_normalized TEXT PRIMARY KEY,
    representative_heading TEXT,
    expected_table_number INTEGER NOT NULL,
    document_count INTEGER NOT NULL,
    occurrence_count INTEGER NOT NULL,
    percent_of_ok_documents REAL NOT NULL,
    style_names TEXT
);
CREATE TABLE IF NOT EXISTS T00_Profile_Section_Coverage (
    expected_table_number INTEGER PRIMARY KEY,
    section_name TEXT NOT NULL,
    expected_container TEXT NOT NULL,
    ok_document_count INTEGER NOT NULL,
    documents_with_section INTEGER NOT NULL,
    occurrence_count INTEGER NOT NULL,
    documents_with_duplicates INTEGER NOT NULL,
    coverage_percent REAL NOT NULL
);
CREATE TABLE IF NOT EXISTS T00_Profile_Exceptions (
    exception_id INTEGER PRIMARY KEY AUTOINCREMENT,
    document_id INTEGER NOT NULL,
    exception_type TEXT NOT NULL,
    expected_table_number INTEGER NOT NULL DEFAULT 0,
    detail TEXT,
    FOREIGN KEY(document_id) REFERENCES T00_Profile_Documents(document_id) ON DELETE CASCADE
);
CREATE TABLE IF NOT EXISTS T00_Profile_Corpus_Summary (
    metric TEXT PRIMARY KEY,
    value_text TEXT,
    value_number REAL
);

CREATE INDEX IF NOT EXISTS IX_Profile_Documents_Status ON T00_Profile_Documents(status, is_present);
CREATE INDEX IF NOT EXISTS IX_Profile_Documents_Layout ON T00_Profile_Documents(layout_signature);
CREATE INDEX IF NOT EXISTS IX_Profile_Tables_Document ON T00_Profile_Table_Occurrences(document_id);
CREATE INDEX IF NOT EXISTS IX_Profile_Tables_Section ON T00_Profile_Table_Occurrences(expected_table_number);
CREATE INDEX IF NOT EXISTS IX_Profile_Tables_Signature ON T00_Profile_Table_Occurrences(table_signature);
CREATE INDEX IF NOT EXISTS IX_Profile_Headings_Document ON T00_Profile_Heading_Occurrences(document_id);
CREATE INDEX IF NOT EXISTS IX_Profile_Exceptions_Document ON T00_Profile_Exceptions(document_id);
"""

VIEW_SQL = r"""
DROP VIEW IF EXISTS V00_Corpus_Overview;
CREATE VIEW V00_Corpus_Overview AS SELECT metric, value_text, value_number FROM T00_Profile_Corpus_Summary;
DROP VIEW IF EXISTS V00_Layout_Families;
CREATE VIEW V00_Layout_Families AS SELECT * FROM T00_Profile_Layouts ORDER BY document_count DESC, layout_signature;
DROP VIEW IF EXISTS V00_Table_Signatures;
CREATE VIEW V00_Table_Signatures AS SELECT * FROM T00_Profile_Table_Signatures ORDER BY expected_table_number, document_count DESC;
DROP VIEW IF EXISTS V00_Table_Variants;
CREATE VIEW V00_Table_Variants AS SELECT * FROM T00_Profile_Table_Variants ORDER BY expected_table_number, document_count DESC;
DROP VIEW IF EXISTS V00_Section_Coverage;
CREATE VIEW V00_Section_Coverage AS SELECT * FROM T00_Profile_Section_Coverage ORDER BY expected_table_number;
DROP VIEW IF EXISTS V00_Exception_Worklist;
CREATE VIEW V00_Exception_Worklist AS
SELECT e.exception_id, e.exception_type, e.expected_table_number, e.detail,
       d.relative_path, d.filename, d.status, d.total_table_count, d.known_sections_found
FROM T00_Profile_Exceptions e JOIN T00_Profile_Documents d ON d.document_id=e.document_id
WHERE d.is_present=1 ORDER BY e.exception_type, d.relative_path;
"""


def initialize_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(SCHEMA_SQL)
    conn.executescript(VIEW_SQL)
    conn.commit()


def start_run(conn: sqlite3.Connection, root: Path, db_path: Path, workers: int, full_rescan: bool, files_discovered: int) -> int:
    cur = conn.execute(
        """INSERT INTO T00_Profile_Runs
        (started_at, root_path, output_db, profiler_version, schema_version, python_version,
         platform, workers, full_rescan, files_discovered)
        VALUES (?,?,?,?,?,?,?,?,?,?)""",
        (utc_now(), str(root), str(db_path), PROFILER_VERSION, SCHEMA_VERSION,
         sys.version.split()[0], platform.platform(), workers, int(full_rescan), files_discovered),
    )
    conn.commit()
    return int(cur.lastrowid)


def existing_document_index(conn: sqlite3.Connection) -> dict[str, dict[str, Any]]:
    rows = q_all(conn, "SELECT document_id, relative_path, file_size_bytes, modified_ns, status FROM T00_Profile_Documents")
    return {str(r['relative_path']): r for r in rows}


def mark_seen_skipped(conn: sqlite3.Connection, relative_path: str, run_id: int) -> None:
    conn.execute("UPDATE T00_Profile_Documents SET is_present=1, last_seen_run_id=? WHERE relative_path=?", (run_id, relative_path))


def upsert_profile_result(conn: sqlite3.Connection, run_id: int, result: dict[str, Any]) -> None:
    now = utc_now()
    rel = result['relative_path']
    existing = q_one(conn, "SELECT document_id FROM T00_Profile_Documents WHERE document_key=?", (result['document_key'],))
    cols = (
        'document_key','relative_path','filename','parent_folder','file_size_bytes','modified_ns','sha256',
        'profiled_at','status','error_type','error_message','paragraph_count','heading_count','body_table_count',
        'header_table_count','footer_table_count','total_table_count','known_sections_found','unknown_heading_count',
        'layout_signature','layout_sequence','variant_layout_signature','variant_layout_sequence','section_number_sequence',
        'elapsed_seconds','is_present','last_seen_run_id'
    )
    vals = (
        result['document_key'], rel, result['filename'], result['parent_folder'], result['file_size_bytes'], result['modified_ns'],
        result.get('sha256',''), now, result['status'], result.get('error_type',''), result.get('error_message',''),
        result.get('paragraph_count',0), result.get('heading_count',0), result.get('body_table_count',0),
        result.get('header_table_count',0), result.get('footer_table_count',0), result.get('total_table_count',0),
        result.get('known_sections_found',0), result.get('unknown_heading_count',0), result.get('layout_signature',''),
        result.get('layout_sequence','[]'), result.get('variant_layout_signature',''), result.get('variant_layout_sequence','[]'),
        result.get('section_number_sequence',''), result.get('elapsed_seconds',0.0), 1, run_id
    )
    if existing:
        document_id = int(existing['document_id'])
        sets = ','.join(f"{c}=?" for c in cols)
        conn.execute(f"UPDATE T00_Profile_Documents SET {sets} WHERE document_id=?", vals + (document_id,))
        conn.execute("DELETE FROM T00_Profile_Heading_Occurrences WHERE document_id=?", (document_id,))
        conn.execute("DELETE FROM T00_Profile_Table_Occurrences WHERE document_id=?", (document_id,))
        conn.execute("DELETE FROM T00_Profile_Exceptions WHERE document_id=?", (document_id,))
    else:
        placeholders = ','.join('?' for _ in cols)
        cur = conn.execute(f"INSERT INTO T00_Profile_Documents ({','.join(cols)}) VALUES ({placeholders})", vals)
        document_id = int(cur.lastrowid)

    if result['status'] == 'ok':
        conn.executemany(
            """INSERT INTO T00_Profile_Heading_Occurrences
            (document_id,heading_ordinal,heading_text,heading_normalized,style_name,expected_table_number)
            VALUES (?,?,?,?,?,?)""",
            [(document_id,i,h['heading_text'],h['heading_normalized'],h['style_name'],h['expected_table_number'])
             for i,h in enumerate(result.get('headings',[]),1)]
        )
        conn.executemany(
            """INSERT INTO T00_Profile_Table_Occurrences
            (document_id,source_order,container_type,container_index,table_ordinal,expected_table_number,detected_section,
             preceding_heading,heading_style,table_signature,table_variant_signature,row_count,min_grid_columns,max_grid_columns,
             row_width_pattern,physical_row_width_pattern,first_row_merge_pattern,generic_token_pattern,physical_cell_count,
             logical_cell_count,horizontal_merge_count,vertical_merge_count,nested_table_count,cell_shape,generic_label_tokens)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            [(document_id,t['source_order'],t['container_type'],t['container_index'],t['table_ordinal'],t['expected_table_number'],
              t['detected_section'],t['preceding_heading'],t['heading_style'],t['table_signature'],t['table_variant_signature'],
              t['row_count'],t['min_grid_columns'],t['max_grid_columns'],t['row_width_pattern'],t['physical_row_width_pattern'],
              t['first_row_merge_pattern'],t['generic_token_pattern'],t['physical_cell_count'],t['logical_cell_count'],
              t['horizontal_merge_count'],t['vertical_merge_count'],t['nested_table_count'],t['cell_shape'],t['generic_label_tokens'])
             for t in result.get('tables',[])]
        )
    else:
        conn.execute(
            """INSERT INTO T00_Profile_Run_Errors
            (run_id,document_key,relative_path,error_type,error_message,traceback,created_at)
            VALUES (?,?,?,?,?,?,?)""",
            (run_id,result['document_key'],rel,result.get('error_type',''),result.get('error_message',''),result.get('traceback',''),now)
        )


def rebuild_aggregates(conn: sqlite3.Connection) -> None:
    for name in ('T00_Profile_Layouts','T00_Profile_Variant_Layouts','T00_Profile_Table_Signatures','T00_Profile_Table_Variants',
                 'T00_Profile_Heading_Variants','T00_Profile_Section_Coverage','T00_Profile_Exceptions','T00_Profile_Corpus_Summary'):
        conn.execute(f"DELETE FROM {name}")

    ok_docs = int(scalar(conn, "SELECT COUNT(*) FROM T00_Profile_Documents WHERE is_present=1 AND status='ok'") or 0)
    denom = max(ok_docs, 1)

    conn.execute("""INSERT INTO T00_Profile_Layouts
        SELECT layout_signature, COUNT(*), ROUND(COUNT(*)*100.0/?,4),
               MIN(section_number_sequence), MIN(layout_sequence), MIN(total_table_count), MAX(total_table_count), AVG(total_table_count),
               MIN(body_table_count), MAX(body_table_count), MIN(header_table_count), MAX(header_table_count), MIN(document_key)
        FROM T00_Profile_Documents WHERE is_present=1 AND status='ok'
        GROUP BY layout_signature""", (denom,))

    conn.execute("""INSERT INTO T00_Profile_Variant_Layouts
        SELECT variant_layout_signature, COUNT(*), ROUND(COUNT(*)*100.0/?,4),
               MIN(section_number_sequence), MIN(variant_layout_sequence), MIN(document_key)
        FROM T00_Profile_Documents WHERE is_present=1 AND status='ok'
        GROUP BY variant_layout_signature""", (denom,))

    conn.execute("""INSERT INTO T00_Profile_Table_Signatures
        SELECT t.table_signature, MIN(t.expected_table_number), MIN(t.detected_section), MIN(t.container_type),
               COUNT(DISTINCT t.document_id), COUNT(*), ROUND(COUNT(DISTINCT t.document_id)*100.0/?,4),
               MIN(t.row_count), MAX(t.row_count), AVG(t.row_count), MIN(t.min_grid_columns), MAX(t.max_grid_columns),
               MIN(t.row_width_pattern), MIN(t.physical_row_width_pattern), MIN(t.horizontal_merge_count), MIN(t.vertical_merge_count),
               MIN(t.nested_table_count), MIN(t.cell_shape), MIN(t.generic_label_tokens)
        FROM T00_Profile_Table_Occurrences t JOIN T00_Profile_Documents d ON d.document_id=t.document_id
        WHERE d.is_present=1 AND d.status='ok' GROUP BY t.table_signature""", (denom,))

    conn.execute("""INSERT INTO T00_Profile_Table_Variants
        SELECT t.table_variant_signature, MIN(t.table_signature), MIN(t.expected_table_number), MIN(t.detected_section), MIN(t.container_type),
               COUNT(DISTINCT t.document_id), COUNT(*), ROUND(COUNT(DISTINCT t.document_id)*100.0/?,4),
               MIN(t.row_count), MAX(t.row_count), AVG(t.row_count), MIN(t.min_grid_columns), MAX(t.max_grid_columns),
               MIN(t.first_row_merge_pattern), MIN(t.generic_token_pattern), MIN(t.generic_label_tokens)
        FROM T00_Profile_Table_Occurrences t JOIN T00_Profile_Documents d ON d.document_id=t.document_id
        WHERE d.is_present=1 AND d.status='ok' GROUP BY t.table_variant_signature""", (denom,))

    conn.execute("""INSERT INTO T00_Profile_Heading_Variants
        SELECT h.heading_normalized, MIN(h.heading_text), MIN(h.expected_table_number), COUNT(DISTINCT h.document_id), COUNT(*),
               ROUND(COUNT(DISTINCT h.document_id)*100.0/?,4), GROUP_CONCAT(DISTINCT h.style_name)
        FROM T00_Profile_Heading_Occurrences h JOIN T00_Profile_Documents d ON d.document_id=h.document_id
        WHERE d.is_present=1 AND d.status='ok' GROUP BY h.heading_normalized""", (denom,))

    for n, meta in EXPECTED_TABLES.items():
        docs_with = int(scalar(conn, """SELECT COUNT(DISTINCT t.document_id) FROM T00_Profile_Table_Occurrences t
            JOIN T00_Profile_Documents d ON d.document_id=t.document_id
            WHERE d.is_present=1 AND d.status='ok' AND t.expected_table_number=?""", (n,)) or 0)
        occ = int(scalar(conn, """SELECT COUNT(*) FROM T00_Profile_Table_Occurrences t JOIN T00_Profile_Documents d ON d.document_id=t.document_id
            WHERE d.is_present=1 AND d.status='ok' AND t.expected_table_number=?""", (n,)) or 0)
        dup = int(scalar(conn, """SELECT COUNT(*) FROM (SELECT t.document_id FROM T00_Profile_Table_Occurrences t
            JOIN T00_Profile_Documents d ON d.document_id=t.document_id WHERE d.is_present=1 AND d.status='ok' AND t.expected_table_number=?
            GROUP BY t.document_id HAVING COUNT(*)>1)""", (n,)) or 0)
        conn.execute("INSERT INTO T00_Profile_Section_Coverage VALUES (?,?,?,?,?,?,?,?)",
                     (n,meta['section'],meta['container'],ok_docs,docs_with,occ,dup,round(docs_with*100.0/denom,4)))

    # Exception worklist
    docs = q_all(conn, "SELECT * FROM T00_Profile_Documents WHERE is_present=1")
    for d in docs:
        did = d['document_id']
        if d['status'] != 'ok':
            conn.execute("INSERT INTO T00_Profile_Exceptions(document_id,exception_type,detail) VALUES (?,?,?)",
                         (did,'parse_error',f"{d.get('error_type','')}: {d.get('error_message','')}"))
            continue
        if int(d.get('total_table_count') or 0) != 11:
            conn.execute("INSERT INTO T00_Profile_Exceptions(document_id,exception_type,detail) VALUES (?,?,?)",
                         (did,'table_count_not_11',f"Found {d.get('total_table_count')} total tables"))
        counts = {int(r['expected_table_number']): int(r['cnt']) for r in q_all(conn,
            "SELECT expected_table_number,COUNT(*) cnt FROM T00_Profile_Table_Occurrences WHERE document_id=? GROUP BY expected_table_number", (did,))}
        for n in EXPECTED_TABLES:
            c = counts.get(n,0)
            if c == 0:
                conn.execute("INSERT INTO T00_Profile_Exceptions(document_id,exception_type,expected_table_number,detail) VALUES (?,?,?,?)",
                             (did,'missing_expected_table',n,EXPECTED_TABLES[n]['section']))
            elif c > 1:
                conn.execute("INSERT INTO T00_Profile_Exceptions(document_id,exception_type,expected_table_number,detail) VALUES (?,?,?,?)",
                             (did,'duplicate_expected_table',n,f"{EXPECTED_TABLES[n]['section']}: {c} occurrences"))
        unmapped = counts.get(0,0)
        if unmapped:
            conn.execute("INSERT INTO T00_Profile_Exceptions(document_id,exception_type,detail) VALUES (?,?,?)",
                         (did,'unmapped_table',f"{unmapped} table(s) were not mapped to T01-T11"))
        if int(d.get('unknown_heading_count') or 0):
            conn.execute("INSERT INTO T00_Profile_Exceptions(document_id,exception_type,detail) VALUES (?,?,?)",
                         (did,'unknown_heading',f"{d.get('unknown_heading_count')} heading(s) were not recognized"))

    metrics = {
        'documents_present': int(scalar(conn,"SELECT COUNT(*) FROM T00_Profile_Documents WHERE is_present=1") or 0),
        'documents_ok': ok_docs,
        'documents_failed': int(scalar(conn,"SELECT COUNT(*) FROM T00_Profile_Documents WHERE is_present=1 AND status<>'ok'") or 0),
        'success_percent': round(ok_docs*100.0/max(int(scalar(conn,"SELECT COUNT(*) FROM T00_Profile_Documents WHERE is_present=1") or 0),1),4),
        'documents_exactly_11_tables': int(scalar(conn,"SELECT COUNT(*) FROM T00_Profile_Documents WHERE is_present=1 AND status='ok' AND total_table_count=11") or 0),
        'documents_all_11_sections_recognized': int(scalar(conn,"SELECT COUNT(*) FROM T00_Profile_Documents WHERE is_present=1 AND status='ok' AND known_sections_found=11") or 0),
        'documents_with_exceptions': int(scalar(conn,"SELECT COUNT(DISTINCT document_id) FROM T00_Profile_Exceptions") or 0),
        'unique_layout_families': int(scalar(conn,"SELECT COUNT(*) FROM T00_Profile_Layouts") or 0),
        'unique_variant_layouts': int(scalar(conn,"SELECT COUNT(*) FROM T00_Profile_Variant_Layouts") or 0),
        'unique_table_signatures': int(scalar(conn,"SELECT COUNT(*) FROM T00_Profile_Table_Signatures") or 0),
        'unique_table_variants': int(scalar(conn,"SELECT COUNT(*) FROM T00_Profile_Table_Variants") or 0),
        'unmapped_table_occurrences': int(scalar(conn,"""SELECT COUNT(*) FROM T00_Profile_Table_Occurrences t JOIN T00_Profile_Documents d ON d.document_id=t.document_id
              WHERE d.is_present=1 AND d.status='ok' AND t.expected_table_number=0""") or 0),
    }
    for k,v in metrics.items():
        conn.execute("INSERT INTO T00_Profile_Corpus_Summary(metric,value_text,value_number) VALUES (?,?,?)", (k,str(v),float(v)))
    conn.commit()


def export_csv(conn: sqlite3.Connection, sql: str, path: Path) -> None:
    cur = conn.execute(sql)
    headers = [d[0] for d in cur.description]
    with path.open('w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(cur.fetchall())


def export_aggregate_outputs(conn: sqlite3.Connection, output_dir: Path) -> list[Path]:
    outputs: list[Path] = []
    jobs = {
        'corpus_summary.csv': "SELECT metric,value_text,value_number FROM T00_Profile_Corpus_Summary ORDER BY metric",
        'layout_families.csv': "SELECT * FROM T00_Profile_Layouts ORDER BY document_count DESC,layout_signature",
        'variant_layouts.csv': "SELECT * FROM T00_Profile_Variant_Layouts ORDER BY document_count DESC,variant_layout_signature",
        'table_signatures.csv': "SELECT * FROM T00_Profile_Table_Signatures ORDER BY expected_table_number,document_count DESC",
        'table_variants.csv': "SELECT * FROM T00_Profile_Table_Variants ORDER BY expected_table_number,document_count DESC",
        'section_coverage.csv': "SELECT * FROM T00_Profile_Section_Coverage ORDER BY expected_table_number",
        'heading_variants.csv': "SELECT * FROM T00_Profile_Heading_Variants ORDER BY expected_table_number,document_count DESC",
        'exception_summary.csv': "SELECT exception_type,COUNT(*) occurrence_count,COUNT(DISTINCT document_id) document_count FROM T00_Profile_Exceptions GROUP BY exception_type ORDER BY document_count DESC,exception_type",
    }
    for name,sql in jobs.items():
        p=output_dir/name; export_csv(conn,sql,p); outputs.append(p)
    summary = {r['metric']: r['value_number'] for r in q_all(conn,"SELECT metric,value_number FROM T00_Profile_Corpus_Summary")}
    jp=output_dir/'corpus_profile_summary.json'
    jp.write_text(json.dumps({'profiler_version':PROFILER_VERSION,'schema_version':SCHEMA_VERSION,'generated_at':utc_now(),'summary':summary},indent=2),encoding='utf-8')
    outputs.append(jp)
    zp=output_dir/'spec_corpus_profile_shareable.zip'
    with zipfile.ZipFile(zp,'w',zipfile.ZIP_DEFLATED) as z:
        for p in outputs:
            z.write(p,arcname=p.name)
    outputs.append(zp)
    return outputs


def print_summary(conn: sqlite3.Connection) -> None:
    print("\nCORPUS SUMMARY")
    print("-"*60)
    for r in q_all(conn,"SELECT metric,value_text FROM T00_Profile_Corpus_Summary ORDER BY metric"):
        print(f"{r['metric']:<42} {r['value_text']}")


def backup_and_remove_db(db_path: Path) -> Path | None:
    if not db_path.exists():
        return None
    stamp=datetime.now().strftime('%Y%m%d_%H%M%S')
    backup=db_path.with_name(f"{db_path.stem}_backup_{stamp}{db_path.suffix}")
    shutil.copy2(db_path,backup)
    # Remove SQLite WAL companions before replacing main database.
    for side in (Path(str(db_path)+'-wal'),Path(str(db_path)+'-shm')):
        if side.exists():
            side.unlink()
    db_path.unlink()
    return backup


def parse_args() -> argparse.Namespace:
    p=argparse.ArgumentParser(description="Recursively profile .docx specs under the current Windows user's Desktop\\Specs folder into SQLite.")
    p.add_argument('--root',type=Path,default=None,help="Override source root. Default: current Windows user's Desktop\\Specs.")
    p.add_argument('--output',type=Path,default=None,help="Output folder. Default: <Specs>\\_Corpus_Profile.")
    p.add_argument('--db',type=Path,default=None,help=f"SQLite profile database path. Default: <output>\\{DEFAULT_DB_FILENAME}.")
    p.add_argument('--workers',type=int,default=min(4,max(1,(os.cpu_count() or 2)-1)),help="Parallel parser processes. Default: max 4.")
    p.add_argument('--full-rescan',action='store_true',help="Reparse unchanged documents.")
    p.add_argument('--rebuild',action='store_true',help="Back up and rebuild the SQLite profile database.")
    p.add_argument('--exclude-dir',action='append',default=[],help="Directory name to skip; may be supplied multiple times.")
    p.add_argument('--progress-every',type=int,default=100,help="Print progress every N completed files.")
    p.add_argument('--verbose',action='store_true',help="Print one line per parsed document.")
    p.add_argument('--no-exports',action='store_true',help="Do not create aggregate CSV/JSON/shareable ZIP outputs.")
    return p.parse_args()


def main() -> int:
    args=parse_args()
    root=(args.root.expanduser().resolve() if args.root is not None else default_specs_root().resolve())
    if not root.exists() or not root.is_dir():
        print(f"ERROR: Specs folder does not exist: {root}\nCreate a folder named 'Specs' on the current user's Windows Desktop, or pass --root to override it.",file=sys.stderr)
        return 2
    output=(args.output.expanduser().resolve() if args.output is not None else (root/DEFAULT_OUTPUT_DIRNAME).resolve())
    output.mkdir(parents=True,exist_ok=True)
    db_path=(args.db.expanduser().resolve() if args.db is not None else output/DEFAULT_DB_FILENAME)
    if args.rebuild:
        backup=backup_and_remove_db(db_path)
        if backup: print(f"Backup created:    {backup}")
    workers=max(1,int(args.workers)); started=time.perf_counter()
    print("Specification Corpus Profiler - SQLite")
    print(f"Profiler version: {PROFILER_VERSION}")
    print(f"Specs folder:     {root}")
    print(f"Output folder:    {output}")
    print(f"SQLite database:  {db_path}")
    print(f"Workers:          {workers}")
    print("Recursive scan:   YES")
    conn=connect_db(db_path); initialize_schema(conn)
    discovered=discover_docx(root,output,set(args.exclude_dir))
    print(f"Word files found: {len(discovered):,}")
    if not discovered:
        conn.close(); return 0
    run_id=start_run(conn,root,db_path,workers,args.full_rescan,len(discovered))
    existing=existing_document_index(conn); to_profile=[]; skipped=0
    for path,rel,size,mtime in discovered:
        old=existing.get(rel)
        unchanged=(old is not None and old['status']=='ok' and int(old['file_size_bytes'] if old['file_size_bytes'] is not None else -2)==int(size)
                   and int(old['modified_ns'] if old['modified_ns'] is not None else -2)==int(mtime))
        if unchanged and not args.full_rescan:
            mark_seen_skipped(conn,rel,run_id); skipped+=1
        else: to_profile.append((str(path),rel,size,mtime))
    conn.commit(); print(f"Unchanged/skipped: {skipped:,}"); print(f"To profile:        {len(to_profile):,}")
    profiled=failed=batch=0
    def handle(result):
        nonlocal profiled,failed,batch
        upsert_profile_result(conn,run_id,result); profiled+=1; batch+=1
        if result['status']!='ok':
            failed+=1; print(f"ERROR: {result['relative_path']} :: {result['error_type']}: {result['error_message']}")
        elif args.verbose:
            print(f"OK: {result['relative_path']} :: tables={result['total_table_count']} known={result['known_sections_found']} layout={result['layout_signature']}")
        if batch>=100: conn.commit(); batch=0
        if profiled % max(1,args.progress_every)==0 or profiled==len(to_profile):
            elapsed=max(time.perf_counter()-started,.001); rate=(profiled+skipped)/elapsed; rem=len(to_profile)-profiled; eta=rem/max(rate,.001)
            print(f"Progress: {profiled:,}/{len(to_profile):,} parsed; {skipped:,} cached; {failed:,} failed; {rate:.1f} files/sec; ETA ~{eta/60:.1f} min")
    if to_profile:
        if workers==1:
            for task in to_profile: handle(profile_one_file(task))
        else:
            with ProcessPoolExecutor(max_workers=workers) as pool:
                futures=[pool.submit(profile_one_file,t) for t in to_profile]
                for fut in as_completed(futures):
                    try: handle(fut.result())
                    except Exception as exc:
                        failed+=1; profiled+=1
                        conn.execute("INSERT INTO T00_Profile_Run_Errors(run_id,error_type,error_message,traceback,created_at) VALUES (?,?,?,?,?)",
                                     (run_id,type(exc).__name__,str(exc),traceback.format_exc(limit=8),utc_now()))
                        print(f"WORKER ERROR: {type(exc).__name__}: {exc}")
    conn.commit()
    conn.execute("UPDATE T00_Profile_Documents SET is_present=CASE WHEN last_seen_run_id=? THEN 1 ELSE 0 END",(run_id,)); conn.commit()
    print("Building aggregate profile tables..."); rebuild_aggregates(conn)
    elapsed=time.perf_counter()-started
    conn.execute("UPDATE T00_Profile_Runs SET completed_at=?,files_profiled=?,files_skipped=?,files_failed=?,elapsed_seconds=? WHERE run_id=?",
                 (utc_now(),profiled,skipped,failed,round(elapsed,3),run_id)); conn.commit()
    if not args.no_exports:
        exported=export_aggregate_outputs(conn,output)
        print("\nAggregate exports created:")
        for p in exported: print(f"  {p}")
    print_summary(conn); conn.close()
    print(f"\nCompleted in {elapsed/60:.2f} minutes.")
    print(f"SQLite profile database: {db_path}")
    if not args.no_exports: print(f"Shareable aggregate bundle: {output/'spec_corpus_profile_shareable.zip'}")
    return 0 if failed==0 else 1


if __name__ == '__main__':
    raise SystemExit(main())
